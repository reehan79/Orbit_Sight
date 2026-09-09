"""Write Evaluation_Metrics.xlsx from frozen detections + GT (post-inference only)."""

from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from orbitsight.evaluation.tii_style import load_tii_gt, match_and_score

_BLUE_HEADER = "FF1F3864"
_BLUE_TRAIN = "FFD6E4F0"
_GREEN_TEST = "FFE2EFDA"
_YELLOW_TOTAL = "FFFFF2CC"


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _border() -> Border:
    side = Side(style="thin", color="FF000000")
    return Border(left=side, right=side, top=side, bottom=side)


def _fmt(x: float | None) -> float | str:
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return ""
    return float(f"{float(x):.6f}")


def find_gt_path(search_roots: list[Path], sequence: str) -> Path | None:
    name = f"{sequence}_bb_windows_40ms.txt"
    for root in search_roots:
        if not root:
            continue
        p = Path(root) / name
        if p.is_file():
            return p
    return None


def score_sequence(
    sequence: str,
    internal_rows: list[tuple],
    gt_path: Path,
    *,
    label: str = "",
    iou_thresh: float = 0.5,
) -> dict:
    gt = load_tii_gt(gt_path)
    # match_and_score expects ranked preds with confidence; keep emission order then score.
    preds = [
        (int(ws), int(we), int(round(cx)), int(round(cy)), int(round(w)), int(round(h)), float(conf))
        for ws, we, cx, cy, w, h, conf in internal_rows
    ]
    m = match_and_score(gt, preds, iou_thresh=iou_thresh)
    return {
        "label": label,
        "seq": sequence,
        "n_gt": m.n_gt,
        "n_pred": m.n_pred,
        "prec": m.precision,
        "rec": m.recall,
        "f1": m.f1,
        "ap": m.ap50,
        "tp": m.tp,
        "fp": m.fp,
        "fn": m.fn,
    }


def write_evaluation_metrics_xlsx(
    excel_path: Path,
    results: list[dict],
    *,
    latency: dict | None = None,
) -> Path:
    """Organizer-style workbook; filename must be Evaluation_Metrics.xlsx."""
    excel_path = Path(excel_path)
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation"

    headers = [
        "Type",
        "Sequence",
        "GT Boxes",
        "Pred Boxes",
        "Precision",
        "Recall",
        "F1 Score",
        "AP @ IoU 0.5",
        "TP",
        "FP",
        "FN",
    ]
    header_font = Font(name="Calibri", bold=True, color="FFFFFFFF", size=11)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = _fill(_BLUE_HEADER)
        cell.border = _border()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    all_tp = all_fp = all_fn = 0
    all_ap: list[float] = []
    row_idx = 2
    for r in results:
        fill_color = _BLUE_TRAIN if r.get("label") == "Training" else _GREEN_TEST
        values = [
            r.get("label", ""),
            r["seq"],
            r["n_gt"],
            r["n_pred"],
            _fmt(r["prec"]),
            _fmt(r["rec"]),
            _fmt(r["f1"]),
            _fmt(r["ap"]),
            r["tp"],
            r["fp"],
            r["fn"],
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=v)
            cell.font = Font(name="Calibri", size=11)
            cell.fill = _fill(fill_color)
            cell.border = _border()
            cell.alignment = Alignment(
                horizontal="center" if col != 2 else "left",
                vertical="center",
            )
        all_tp += int(r["tp"])
        all_fp += int(r["fp"])
        all_fn += int(r["fn"])
        if r["ap"] is not None and not (isinstance(r["ap"], float) and np.isnan(r["ap"])):
            all_ap.append(float(r["ap"]))
        row_idx += 1

    overall_prec = all_tp / (all_tp + all_fp) if all_tp + all_fp else 0.0
    overall_rec = all_tp / (all_tp + all_fn) if all_tp + all_fn else 0.0
    overall_f1 = (
        2 * overall_prec * overall_rec / (overall_prec + overall_rec)
        if overall_prec + overall_rec
        else 0.0
    )
    map50 = float(np.mean(all_ap)) if all_ap else float("nan")
    total_values = [
        "Overall",
        f"Average ({len(results)} sequences)",
        "",
        "",
        _fmt(overall_prec),
        _fmt(overall_rec),
        _fmt(overall_f1),
        _fmt(map50),
        all_tp,
        all_fp,
        all_fn,
    ]
    for col, v in enumerate(total_values, 1):
        cell = ws.cell(row=row_idx, column=col, value=v)
        cell.font = Font(name="Calibri", bold=True, size=11)
        cell.fill = _fill(_YELLOW_TOTAL)
        cell.border = _border()
        cell.alignment = Alignment(
            horizontal="center" if col != 2 else "left",
            vertical="center",
        )

    col_widths = [12, 52, 10, 12, 11, 9, 11, 14, 7, 7, 7]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    # Inference efficiency sheet (challenge outcome metrics request).
    ws2 = wb.create_sheet("InferenceEfficiency")
    ws2["A1"] = "Metric"
    ws2["B1"] = "Value"
    lat = latency or {}
    rows_lat = [
        ("sequences", lat.get("n_sequences", "")),
        ("total_wall_s", lat.get("total_wall_s", "")),
        ("mean_seq_wall_s", lat.get("mean_seq_wall_s", "")),
        ("p95_seq_wall_s", lat.get("p95_seq_wall_s", "")),
        ("note", lat.get("note", "Wall-clock per sequence during Docker/native inference")),
    ]
    for i, (k, v) in enumerate(rows_lat, start=2):
        ws2.cell(row=i, column=1, value=k)
        ws2.cell(row=i, column=2, value=v)

    wb.save(excel_path)
    return excel_path
