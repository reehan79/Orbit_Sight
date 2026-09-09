"""Official public-timeline output contract V2 (no Testing_sets / sealed rerun)."""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from orbitsight.submission import (
    infer_sequence,
    load_final_bundle,
    output_dir,
    run_dataset,
    write_sequence_prediction,
)
from orbitsight.submission.public_contract import (
    DEFAULT_CLASS_ID,
    PUBLIC_COLUMNS,
    prediction_filename,
    write_public_prediction_file,
)

MODEL = Path(__file__).resolve().parents[1] / "models" / "final"
TRAIN = Path(
    r"D:\OrbitSight_SSA_Challenge\OrbitSight_SSA_Challenge\Phase_1\OrbitSight_Dataset\Training_sets"
)
SEQ = "DVX_NOAA6_11416_2025-01-20-19-06-31"


def test_prediction_filename_modes():
    assert prediction_filename("SEQ", "admin_pred") == "SEQ_pred.txt"
    assert prediction_filename("SEQ", "public_plain") == "SEQ.txt"


def test_public_writer_nine_fields_and_mappings():
    conf = 0.848222565945547
    rows = [(100, 40100, 11.4, 12.6, 5.2, 6.8, conf)]
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / prediction_filename("SYNTH", "admin_pred")
        write_public_prediction_file(path, "SYNTH", rows, class_id=DEFAULT_CLASS_ID)
        text = path.read_text(encoding="utf-8")
        header = text.splitlines()[0].split("\t")
        assert header == list(PUBLIC_COLUMNS)
        parts = text.splitlines()[1].split("\t")
        assert parts[0] == "SYNTH"
        assert parts[1] == "100"
        assert parts[2] == "40100"
        assert parts[3] == "11"  # centre_x from center_x
        assert parts[4] == "13"  # centre_y
        assert parts[5] == "5"  # w
        assert parts[6] == "7"  # h
        assert parts[7] == str(DEFAULT_CLASS_ID)
        assert abs(float(parts[8]) - conf) <= 1e-12


def test_work_output_path_sparsesight():
    with tempfile.TemporaryDirectory() as td:
        out = output_dir(Path(td), team_name="SparseSight-SSA", day="09092026")
        assert out == Path(td) / "SparseSight-SSA" / "09092026"


@pytest.mark.skipif(not (MODEL / "manifest.json").exists(), reason="models/final missing")
@pytest.mark.skipif(not (TRAIN / f"{SEQ}_labeled_events.npy").exists(), reason="Training fixture missing")
def test_run_dataset_public_contract_and_metrics_training_fixture():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        ds = td / "dataset"
        ds.mkdir()
        for suf in ("_labeled_events.npy", "_bb_windows_40ms.txt"):
            src = TRAIN / f"{SEQ}{suf}"
            (ds / f"{SEQ}{suf}").write_bytes(src.read_bytes())
        work = td / "work"
        out = run_dataset(
            ds,
            work,
            model_dir=MODEL,
            team_name="SparseSight-SSA",
            day="09092026",
            filename_mode="admin_pred",
            class_id=1,
            write_metrics=True,
        )
        pred = out / f"{SEQ}_pred.txt"
        assert pred.exists()
        header = pred.read_text(encoding="utf-8").splitlines()[0].split("\t")
        assert header == list(PUBLIC_COLUMNS)
        xlsx = out / "Evaluation_Metrics.xlsx"
        assert xlsx.exists()
        wb = load_workbook(xlsx)
        assert "Evaluation" in wb.sheetnames
        assert "InferenceEfficiency" in wb.sheetnames


@pytest.mark.skipif(not (MODEL / "manifest.json").exists(), reason="models/final missing")
@pytest.mark.skipif(not (TRAIN / f"{SEQ}_labeled_events.npy").exists(), reason="Training fixture missing")
def test_predictions_identical_when_gt_absent_or_changed():
    """GT may affect metrics only — never prediction txt bytes."""
    bundle = load_final_bundle(MODEL)
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        base_ds = td / "base"
        base_ds.mkdir()
        events = TRAIN / f"{SEQ}_labeled_events.npy"
        gt = TRAIN / f"{SEQ}_bb_windows_40ms.txt"
        (base_ds / events.name).write_bytes(events.read_bytes())
        (base_ds / gt.name).write_bytes(gt.read_bytes())

        work_a = td / "work_a"
        out_a = run_dataset(
            base_ds,
            work_a,
            model_dir=MODEL,
            day="09092026",
            filename_mode="admin_pred",
            write_metrics=True,
        )
        pred_a = (out_a / f"{SEQ}_pred.txt").read_bytes()

        # No GT
        no_gt = td / "no_gt"
        no_gt.mkdir()
        (no_gt / events.name).write_bytes(events.read_bytes())
        work_b = td / "work_b"
        out_b = run_dataset(
            no_gt,
            work_b,
            model_dir=MODEL,
            day="09092026",
            filename_mode="admin_pred",
            write_metrics=True,
        )
        pred_b = (out_b / f"{SEQ}_pred.txt").read_bytes()
        assert pred_a == pred_b
        assert not (out_b / "Evaluation_Metrics.xlsx").exists()
        assert (out_b / "Evaluation_Metrics_UNAVAILABLE.txt").exists()

        # Mutated GT
        mut = td / "mut"
        mut.mkdir()
        (mut / events.name).write_bytes(events.read_bytes())
        (mut / gt.name).write_text(
            "window_start_timestamp_us\twindow_end_timestamp_us\t"
            "center_x\tcenter_y\twidth\theight\n"
            "0\t40000\t1\t1\t2\t2\n",
            encoding="utf-8",
        )
        work_c = td / "work_c"
        out_c = run_dataset(
            mut,
            work_c,
            model_dir=MODEL,
            day="09092026",
            filename_mode="admin_pred",
            write_metrics=True,
        )
        pred_c = (out_c / f"{SEQ}_pred.txt").read_bytes()
        assert pred_a == pred_c

        # Direct infer path also unused GT
        rows = infer_sequence(SEQ, no_gt, bundle, use_fast=True)
        write_sequence_prediction(td / "x", SEQ, rows, filename_mode="admin_pred")
        assert (td / "x" / f"{SEQ}_pred.txt").exists()


def test_class_id_default_is_documented_rso_one():
    assert DEFAULT_CLASS_ID == 1
